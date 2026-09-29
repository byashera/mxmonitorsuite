import os
import io
import json
from typing import Dict, Any, List
from datetime import datetime

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import sys

# PyInstaller içindeki gömülü snap7.dll yolunu zorunlu olarak tanıt
if getattr(sys, 'frozen', False):
    bundle_dir = getattr(sys, '_MEIPASS', os.path.dirname(sys.executable))
    dll_path = os.path.join(bundle_dir, "snap7.dll")
    if os.path.exists(dll_path):
        os.environ["PATH"] = bundle_dir + os.pathsep + os.environ.get("PATH", "")

import snap7
from snap7.util import get_bool, get_int, get_real, get_dword

# Snap7 C-API RemotePort parametre sabiti
PARAM_REMOTE_PORT = 2

router = APIRouter(prefix="/api", tags=["PLC Yönetimi"])

# ==========================================
# 1. DAHİLİ CONFIG & STORAGE YÖNETİCİSİ
# ==========================================
class JSONConfigStore:
    def __init__(self, filepath="storage/plcs.json"):
        self.filepath = filepath
        self._ensure_file()

    def _ensure_file(self):
        os.makedirs(os.path.dirname(self.filepath), exist_ok=True)
        if not os.path.exists(self.filepath):
            with open(self.filepath, "w", encoding="utf-8") as f:
                json.dump({"plcs": [], "watchlists": []}, f, indent=4, ensure_ascii=False)

    def _read_data(self) -> dict:
        try:
            with open(self.filepath, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {"plcs": [], "watchlists": []}

    def _write_data(self, data: dict):
        with open(self.filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)

    def get_all_plcs(self) -> list:
        return self._read_data().get("plcs", [])

    def add_plc(self, plc_data: dict) -> dict:
        data = self._read_data()
        plcs = data.setdefault("plcs", [])
        new_id = max([p.get("id", 0) for p in plcs], default=0) + 1
        plc_data["id"] = new_id
        plc_data.setdefault("port", 102)
        plc_data.setdefault("tags", [])
        plcs.append(plc_data)
        self._write_data(data)
        return plc_data

    def update_plc(self, plc_id: int, plc_data: dict):
        data = self._read_data()
        for p in data.get("plcs", []):
            if p.get("id") == plc_id:
                p["name"] = plc_data.get("name", p["name"])
                p["ip"] = plc_data.get("ip", p["ip"])
                p["port"] = int(plc_data.get("port", p.get("port", 102)))
                p["plant"] = plc_data.get("plant", p.get("plant", "JAN"))
                p["rack"] = plc_data.get("rack", p["rack"])
                p["slot"] = plc_data.get("slot", p["slot"])
                self._write_data(data)
                return p
        raise HTTPException(status_code=404, detail="PLC bulunamadı.")

    def delete_plc(self, plc_id: int):
        data = self._read_data()
        data["plcs"] = [p for p in data.get("plcs", []) if p.get("id") != plc_id]
        self._write_data(data)

    def add_tag_to_plc(self, plc_id: int, tag_data: dict) -> dict:
        data = self._read_data()
        for p in data.get("plcs", []):
            if p.get("id") == plc_id:
                tags = p.setdefault("tags", [])
                new_tag_id = max([t.get("id", 0) for t in tags], default=0) + 1
                tag_data["id"] = new_tag_id
                tags.append(tag_data)
                self._write_data(data)
                return tag_data
        raise HTTPException(status_code=404, detail="PLC bulunamadı.")

    def delete_tag(self, plc_id: int, tag_id: int):
        data = self._read_data()
        for p in data.get("plcs", []):
            if p.get("id") == plc_id:
                p["tags"] = [t for t in p.get("tags", []) if t.get("id") != tag_id]
                self._write_data(data)
                return
        raise HTTPException(status_code=404, detail="Tag bulunamadı.")

    def get_watchlists(self) -> list:
        return self._read_data().get("watchlists", [])

    def save_watchlist(self, list_data: dict) -> dict:
        data = self._read_data()
        lists = data.setdefault("watchlists", [])
        new_id = max([w.get("id", 0) for w in lists], default=0) + 1
        list_data["id"] = new_id
        list_data["created_at"] = datetime.now().strftime("%d.%m.%Y %H:%M")
        lists.append(list_data)
        self._write_data(data)
        return list_data

config_store = JSONConfigStore()


# ==========================================
# 2. PYDANTIC VERİ MODELLERİ
# ==========================================
class TagCreate(BaseModel):
    name: str
    db_number: int
    offset: int
    data_type: str

class TagUpdateSchema(BaseModel):
    name: str
    db_number: int
    offset: int
    data_type: str

class PLCCreate(BaseModel):
    name: str
    ip: str
    port: int = 102
    plant: str = "JAN"
    rack: int = 0
    slot: int = 1
    active: bool = True

class PLCUpdate(BaseModel):
    name: str
    ip: str
    port: int = 102
    plant: str
    rack: int
    slot: int

class ExportRequest(BaseModel):
    plc_id: int
    format: str
    data: Dict[str, Any]

class ConsoleCommand(BaseModel):
    command: str

class WatchlistCreateSchema(BaseModel):
    name: str
    tags: List[Dict[str, Any]]


# ==========================================
# 3. PLC CRUD ROTALARI
# ==========================================
@router.get("/plcs")
async def get_plcs():
    return {"status": "success", "data": config_store.get_all_plcs()}

@router.post("/plcs")
async def create_plc(plc: PLCCreate):
    config_store.add_plc(plc.model_dump())
    return {"status": "success", "message": "Yeni PLC başarıyla eklendi."}

@router.put("/plcs/{plc_id}")
async def update_plc_endpoint(plc_id: int, plc: PLCUpdate):
    config_store.update_plc(plc_id, plc.model_dump())
    return {"status": "success", "message": "PLC başarıyla güncellendi."}

@router.delete("/plcs/{plc_id}")
async def remove_plc(plc_id: int):
    config_store.delete_plc(plc_id)
    return {"status": "success", "message": "PLC silindi."}


# ==========================================
# 4. TAG CRUD ROTALARI
# ==========================================
@router.post("/plcs/{plc_id}/tags")
async def add_tag(plc_id: int, tag: TagCreate):
    config_store.add_tag_to_plc(plc_id, tag.model_dump())
    return {"status": "success", "message": "Tag başarıyla eklendi."}

@router.put("/plcs/{plc_id}/tags/{tag_id}")
async def update_tag(plc_id: int, tag_id: int, tag_data: TagUpdateSchema):
    data = config_store._read_data()
    for plc in data.get("plcs", []):
        if plc.get("id") == plc_id:
            for tag in plc.get("tags", []):
                if tag.get("id") == tag_id:
                    tag["name"] = tag_data.name
                    tag["db_number"] = tag_data.db_number
                    tag["offset"] = tag_data.offset
                    tag["data_type"] = tag_data.data_type
                    config_store._write_data(data)
                    return {"status": "success", "message": "Tag güncellendi."}
    raise HTTPException(status_code=404, detail="Tag veya PLC bulunamadı.")

@router.delete("/plcs/{plc_id}/tags/{tag_id}")
async def remove_tag(plc_id: int, tag_id: int):
    config_store.delete_tag(plc_id, tag_id)
    return {"status": "success", "message": "Tag silindi."}


# ==========================================
# 5. GERÇEK SAHA TEK SEFERLİK SORGULAMA (ONE-SHOT POLL)
# ==========================================
@router.post("/plcs/{plc_id}/poll")
async def poll_single_plc(plc_id: int):
    data = config_store._read_data()
    plc = next((p for p in data.get("plcs", []) if p.get("id") == plc_id), None)
    if not plc:
        raise HTTPException(status_code=404, detail="PLC bulunamadı.")

    poll_time = datetime.now().strftime("%H:%M:%S")
    read_results = []
    
    ip = plc.get("ip")
    port = int(plc.get("port", 102))
    rack = int(plc.get("rack", 0))
    slot = int(plc.get("slot", 1))

    client = snap7.client.Client()
    
    try:
        # Özel port kontrolü (Varsayılan 102'den farklıysa)
        if port != 102:
            client.set_param(PARAM_REMOTE_PORT, port)

        client.connect(ip, rack, slot)
        
        if not client.get_connected():
            return {
                "status": "error",
                "message": f"Bağlantı kurulamadı: {ip}:{port} (Rack: {rack}, Slot: {slot})"
            }

        for tag in plc.get("tags", []):
            db_no = int(tag.get("db_number"))
            offset = int(tag.get("offset"))
            t_type = tag.get("data_type", "Real")
            
            try:
                if t_type == "Bool":
                    raw_data = client.db_read(db_no, offset, 1)
                    val = get_bool(raw_data, 0, 0)
                elif t_type == "Int":
                    raw_data = client.db_read(db_no, offset, 2)
                    val = get_int(raw_data, 0)
                elif t_type == "Real":
                    raw_data = client.db_read(db_no, offset, 4)
                    val = round(get_real(raw_data, 0), 2)
                elif t_type == "DWord":
                    raw_data = client.db_read(db_no, offset, 4)
                    val = get_dword(raw_data, 0)
                else:
                    val = "-"

                read_results.append({
                    "id": tag.get("id"),
                    "name": tag.get("name"),
                    "db_number": db_no,
                    "offset": offset,
                    "type": t_type,
                    "value": val,
                    "status": "Good"
                })
            except Exception:
                read_results.append({
                    "id": tag.get("id"),
                    "name": tag.get("name"),
                    "db_number": db_no,
                    "offset": offset,
                    "type": t_type,
                    "value": "Adres/DB Hatası",
                    "status": "Hata"
                })

    except Exception as e:
        return {
            "status": "error",
            "message": f"Haberleşme hatası: {str(e)}"
        }
    finally:
        try:
            if client.get_connected():
                client.disconnect()
            client.destroy()
        except Exception:
            pass

    return {
        "status": "success",
        "plc_id": plc_id,
        "plc_name": plc.get("name"),
        "plant": plc.get("plant", "JAN"),
        "timestamp": poll_time,
        "tags": read_results
    }


# ==========================================
# 6. İZLEME LİSTELERİ (WATCHLIST) ROTALARI
# ==========================================
@router.get("/watchlists")
async def get_watchlists():
    return {"status": "success", "data": config_store.get_watchlists()}

@router.post("/watchlists")
async def create_watchlist(req: WatchlistCreateSchema):
    saved = config_store.save_watchlist(req.model_dump())
    return {"status": "success", "data": saved, "message": "İzleme listesi oluşturuldu."}

@router.delete("/watchlists/{list_id}")
async def delete_watchlist(list_id: int):
    data = config_store._read_data()
    data["watchlists"] = [w for w in data.get("watchlists", []) if w.get("id") != list_id]
    config_store._write_data(data)
    return {"status": "success", "message": "İzleme listesi silindi."}


# ==========================================
# 7. DIŞARI AKTARMA (EXPORT) ROTASI
# ==========================================
@router.post("/export/data")
async def export_data(req: ExportRequest):
    plc_tags = req.data
    plcs = config_store.get_all_plcs()
    target_plc = next((p for p in plcs if p.get("id") == req.plc_id), None)

    plc_name = target_plc.get("name", f"PLC_{req.plc_id}") if target_plc else f"PLC_{req.plc_id}"
    plc_plant = target_plc.get("plant", "JAN") if target_plc else "JAN"
    plc_ip = target_plc.get("ip", "Bilinmiyor") if target_plc else "Bilinmiyor"
    plc_rack_slot = f"{target_plc.get('rack', 0)} / {target_plc.get('slot', 1)}" if target_plc else "0 / 1"
    report_time = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

    if req.format == "sql":
        sql_lines = [
            "-- ========================================================",
            "-- S7 MONITOR SUITE - ENDÜSTRİYEL VERİ AKTARIMI",
            f"-- Cihaz: {plc_name} | Plant: {plc_plant} | IP: {plc_ip}",
            f"-- Raporlama Tarihi: {report_time}",
            "-- ========================================================",
            "CREATE TABLE IF NOT EXISTS s7_live_telemetry (",
            "    id SERIAL PRIMARY KEY,",
            "    plant VARCHAR(10),",
            "    plc_name VARCHAR(100),",
            "    tag_name VARCHAR(100),",
            "    data_type VARCHAR(20),",
            "    tag_value VARCHAR(50),",
            "    signal_status VARCHAR(30),",
            "    recorded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP",
            ");",
            "BEGIN TRANSACTION;"
        ]
        for tag_name, info in plc_tags.items():
            val = info.get("value", "-") if isinstance(info, dict) else info
            st = info.get("status", "Good") if isinstance(info, dict) else "Good"
            t_type = info.get("type", "Var") if isinstance(info, dict) else "Var"
            sql_lines.append(
                f"INSERT INTO s7_live_telemetry (plant, plc_name, tag_name, data_type, tag_value, signal_status) "
                f"VALUES ('{plc_plant}', '{plc_name}', '{tag_name}', '{t_type}', '{val}', '{st}');"
            )
        sql_lines.append("COMMIT;")

        content = "\n".join(sql_lines).encode("utf-8")
        return StreamingResponse(
            io.BytesIO(content),
            media_type="application/sql",
            headers={"Content-Disposition": f"attachment; filename={plc_plant}_{plc_name}_Veri.sql"}
        )

    elif req.format == "word":
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_after = Pt(4)
        run_title = title_p.add_run("S7 MONITOR SUITE | TELEMETRİ RAPORU")
        run_title.font.name = "Arial"
        run_title.font.size = Pt(18)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(30, 58, 138)

        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(18)
        run_sub = sub_p.add_run(f"Otomasyon ve Veri İzleme Sistemi  •  Rapor Tarihi: {report_time}")
        run_sub.font.size = Pt(9.5)
        run_sub.font.color.rgb = RGBColor(100, 116, 139)

        info_table = doc.add_table(rows=2, cols=2)
        info_table.autofit = False
        info_table.columns[0].width = Inches(3.4)
        info_table.columns[1].width = Inches(3.4)

        panel_data = [
            [f"Tesis / Plant: {plc_plant}", f"PLC Cihazı: {plc_name}"],
            [f"IP Adresi: {plc_ip}", f"Rack / Slot: {plc_rack_slot}"]
        ]

        for r_idx, row in enumerate(panel_data):
            for c_idx, text in enumerate(row):
                cell = info_table.cell(r_idx, c_idx)
                cell.text = text
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                run = p.runs[0]
                run.font.name = "Arial"
                run.font.size = Pt(9.5)
                run.font.bold = True if c_idx == 0 and r_idx == 0 else False
                run.font.color.rgb = RGBColor(15, 23, 42)
                shading = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shading)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        table = doc.add_table(rows=1, cols=4)
        table.autofit = False
        table.columns[0].width = Inches(2.6)
        table.columns[1].width = Inches(1.4)
        table.columns[2].width = Inches(1.5)
        table.columns[3].width = Inches(1.3)

        headers = ["Tag Adı", "Veri Tipi", "Anlık Değer", "Sinyal Durumu"]
        hdr_cells = table.rows[0].cells
        for idx, h in enumerate(headers):
            hdr_cells[idx].text = h
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.runs[0]
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            shading = parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(nsdecls('w')))
            hdr_cells[idx]._tc.get_or_add_tcPr().append(shading)

        for r_num, (tag_name, info) in enumerate(plc_tags.items(), start=1):
            val = str(info.get("value", "-")) if isinstance(info, dict) else str(info)
            st = str(info.get("status", "Good")) if isinstance(info, dict) else "Good"
            t_type = str(info.get("type", "Real")) if isinstance(info, dict) else "Real"
            is_err = "Hata" in val or "Hata" in st or "ulaşılamadı" in val

            row_cells = table.add_row().cells
            row_cells[0].text = tag_name
            row_cells[1].text = t_type
            row_cells[2].text = val
            row_cells[3].text = "Hata" if is_err else "Okunuyor"

            bg_color = "F8FAFC" if r_num % 2 == 0 else "FFFFFF"
            for c_idx, cell in enumerate(row_cells):
                p = cell.paragraphs[0]
                run = p.runs[0]
                run.font.name = "Arial"
                run.font.size = Pt(9.5)

                if c_idx == 3:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(220, 38, 38) if is_err else RGBColor(16, 185, 129)
                elif c_idx == 2 and is_err:
                    run.font.color.rgb = RGBColor(220, 38, 38)
                else:
                    run.font.color.rgb = RGBColor(15, 23, 42)

                shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_color))
                cell._tc.get_or_add_tcPr().append(shading)

        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={plc_plant}_{plc_name}_Rapor.docx"}
        )

    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{plc_plant} - Veri Akışı"
        ws.views.sheetView[0].showGridLines = True

        font_title = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        font_meta = Font(name="Segoe UI", size=9, bold=False, color="475569")
        font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        font_body = Font(name="Segoe UI", size=9.5, color="0F172A")
        font_good = Font(name="Segoe UI", size=9.5, bold=True, color="059669")
        font_error = Font(name="Segoe UI", size=9.5, bold=True, color="DC2626")

        fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_meta = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.merge_cells("A1:E1")
        ws["A1"] = "S7 MONITOR SUITE  |  CANLI VERİ VE DENETİM RAPORU"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 28

        meta_items = [
            ("A2", f"Tesis / Plant: {plc_plant}"),
            ("B2", f"Cihaz: {plc_name}"),
            ("C2", f"IP: {plc_ip}"),
            ("D2", f"Rack/Slot: {plc_rack_slot}"),
            ("E2", f"Tarih: {report_time}")
        ]
        for cell_coord, text in meta_items:
            ws[cell_coord] = text
            ws[cell_coord].font = font_meta
            ws[cell_coord].fill = fill_meta
            ws[cell_coord].alignment = Alignment(vertical="center", horizontal="left")
            ws[cell_coord].border = thin_border
        ws.row_dimensions[2].height = 20

        ws.row_dimensions[3].height = 10

        headers = ["TESİS", "TAG ADI", "VERİ TİPİ", "ANLIK DEĞER", "DURUM"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(vertical="center", horizontal="left" if col_idx <= 2 else "center")
            cell.border = thin_border
        ws.row_dimensions[4].height = 24

        row_num = 5
        for tag_name, info in plc_tags.items():
            val = info.get("value", "-") if isinstance(info, dict) else info
            st = info.get("status", "Good") if isinstance(info, dict) else "Good"
            t_type = info.get("type", "Real") if isinstance(info, dict) else "Real"
            is_err = "Hata" in str(val) or "Hata" in str(st) or "ulaşılamadı" in str(val)

            row_data = [plc_plant, tag_name, t_type, str(val), "HATA" if is_err else "OKUNUYOR"]
            current_fill = fill_zebra if row_num % 2 == 0 else fill_white

            for col_idx, val_item in enumerate(row_data, start=1):
                c = ws.cell(row=row_num, column=col_idx, value=val_item)
                c.fill = current_fill
                c.border = thin_border

                if col_idx == 5:
                    c.font = font_error if is_err else font_good
                    c.alignment = Alignment(vertical="center", horizontal="center")
                elif col_idx in [1, 3, 4]:
                    c.font = font_body
                    c.alignment = Alignment(vertical="center", horizontal="center")
                else:
                    c.font = font_body
                    c.alignment = Alignment(vertical="center", horizontal="left")

            ws.row_dimensions[row_num].height = 20
            row_num += 1

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 1 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={plc_plant}_{plc_name}_Veri.xlsx"}
        )


# ==========================================
# 8. KONSOL ROTASI (PORT DESTEKLİ)
# ==========================================
@router.post("/console")
async def execute_console_command(req: ConsoleCommand):
    cmd_line = req.command.strip()
    if not cmd_line:
        return {"status": "error", "output": "Boş komut girildi."}

    args = cmd_line.split()
    main_cmd = args[0].lower()

    try:
        if main_cmd == "status":
            plcs = config_store._read_data().get("plcs", [])
            if not plcs:
                return {"status": "system", "output": "Sistemde kayıtlı PLC bulunmuyor."}

            out = "Kayıtlı PLC Cihazları:\n"
            for p in plcs:
                state = "AKTİF" if p.get("active") else "PASİF"
                p_port = p.get("port", 102)
                out += f"  [{state}] {p.get('name')} | IP: {p.get('ip')}:{p_port} | Rack:{p.get('rack')} Slot:{p.get('slot')}\n"
            return {"status": "success", "output": out}

        elif main_cmd == "connect":
            if len(args) < 2:
                return {"status": "error", "output": "Kullanım hatası! Örnek: connect 192.168.0.10 veya connect 192.168.0.10:10102"}

            raw_target = args[1]
            if ":" in raw_target:
                ip, port_str = raw_target.split(":", 1)
                port = int(port_str)
            else:
                ip = raw_target
                port = 102

            out = f"Bağlantı deneniyor: {ip}:{port} ...\n"
            client = snap7.client.Client()
            try:
                if port != 102:
                    client.set_param(PARAM_REMOTE_PORT, port)

                client.connect(ip, 0, 1)
                if client.get_connected():
                    out += f"BAŞARILI: {ip}:{port} adresindeki PLC ile el sıkışıldı (S7 Oturumu açıldı)."
                    client.disconnect()
                    client.destroy()
                    return {"status": "success", "output": out}
                else:
                    client.destroy()
                    return {"status": "error", "output": f"BAŞARISIZ: {ip}:{port} adresine ulaşılamadı."}
            except Exception as e:
                return {"status": "error", "output": f"Hata: {str(e)}"}

        elif main_cmd == "read":
            if len(args) < 2:
                return {"status": "error", "output": "Kullanım hatası! Örnek: read DB1.4"}
            return {"status": "system", "output": f"Sorgulanıyor: {args[1]}..."}

        else:
            return {"status": "error", "output": f"Bilinmeyen komut: '{main_cmd}'. Komut listesi için 'help' yazın."}

    except Exception as e:
        return {"status": "error", "output": f"İşlem sırasında beklenmeyen hata oluştu: {str(e)}"}